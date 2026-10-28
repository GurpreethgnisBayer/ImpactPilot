"""PPTX impact brief: 16:9 deck rendered from the shared brief content."""

from __future__ import annotations

import io

from impactpilot.domain.models import InitiativeRecord
from impactpilot.export.brief_content import BriefContent, build_brief_content

INK = (0x0B, 0x12, 0x20)
INDIGO = (0x4F, 0x46, 0xE5)
VIOLET = (0x7C, 0x3A, 0xED)
SLATE = (0x47, 0x55, 0x69)
MUTED = (0x94, 0xA3, 0xB8)
SURFACE = (0xF7, 0xF8, 0xFB)
WHITE = (0xFF, 0xFF, 0xFF)
AMBER = (0xB4, 0x53, 0x09)
AMBER_BG = (0xFE, 0xF3, 0xC7)
GREEN = (0x04, 0x78, 0x57)
FONT = "Arial"


def build_pptx_bytes(record: InitiativeRecord) -> bytes:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("PPTX export requires python-pptx") from exc

    c = build_brief_content(record)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    W = prs.slide_width

    def rgb(t: tuple[int, int, int]) -> RGBColor:
        return RGBColor(*t)

    def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
        s = slide.shapes.add_shape(shape, x, y, w, h)
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
        s.line.fill.background()
        s.shadow.inherit = False
        return s

    def text(slide, x, y, w, h, value, size=14, color=INK, bold=False, align=None):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = value
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color)
        p.font.name = FONT
        if align:
            p.alignment = align
        return tf

    def bullets(tf, items, size=13, color=SLATE):
        for item in items:
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(size)
            p.font.color.rgb = rgb(color)
            p.font.name = FONT
            p.space_before = Pt(6)

    def header(slide, eyebrow: str, title: str):
        rect(slide, 0, 0, W, Inches(1.25), INK)
        rect(slide, 0, Inches(1.25), W, Inches(0.06), INDIGO)
        text(slide, Inches(0.6), Inches(0.2), Inches(9), Inches(0.35), eyebrow.upper(), 11, MUTED, True)
        text(slide, Inches(0.6), Inches(0.5), Inches(10.5), Inches(0.7), title, 26, WHITE, True)
        if c.status == "draft":
            tag = rect(slide, Inches(11.3), Inches(0.42), Inches(1.5), Inches(0.42), AMBER_BG, MSO_SHAPE.ROUNDED_RECTANGLE)
            tf = tag.text_frame
            tf.paragraphs[0].text = "DRAFT"
            tf.paragraphs[0].font.size = Pt(12)
            tf.paragraphs[0].font.bold = True
            tf.paragraphs[0].font.color.rgb = rgb(AMBER)
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        footer = "ImpactPilot · " + ("Draft: LLM-suggested values pending confirmation" if c.status == "draft" else "Confirmed brief")
        text(slide, Inches(0.6), Inches(7.0), Inches(12), Inches(0.3), footer, 9, MUTED)

    _summary_slide(prs, blank, c, header, rect, text, bullets, Inches, Pt, rgb)
    _effects_slide(prs, blank, c, header, text, bullets, Inches, Pt, rgb)
    if c.summary or c.findings or c.risks:
        slide = prs.slides.add_slide(blank)
        header(slide, "Evidence", "What the literature says")
        tf = text(slide, Inches(0.6), Inches(1.6), Inches(7.4), Inches(5.2), c.summary or "", 14, INK)
        bullets(tf, c.findings[:6])
        rect(slide, Inches(8.4), Inches(1.6), Inches(4.3), Inches(5.2), SURFACE, MSO_SHAPE.ROUNDED_RECTANGLE)
        tf2 = text(slide, Inches(8.65), Inches(1.8), Inches(3.9), Inches(4.8), "Risks and caveats", 14, INDIGO, True)
        bullets(tf2, c.risks[:5] or ["None recorded"], 12)
        if c.open_questions:
            p = tf2.add_paragraph()
            p.text = "Open questions"
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = rgb(INDIGO)
            p.space_before = Pt(14)
            bullets(tf2, c.open_questions[:4], 12)
    if c.references:
        slide = prs.slides.add_slide(blank)
        header(slide, "Sources", f"References ({len(c.references)})")
        tf = text(slide, Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.2), "", 11)
        for i, ref in enumerate(c.references[:14], 1):
            p = tf.add_paragraph() if i > 1 else tf.paragraphs[0]
            p.text = f"{i}.  {ref}"
            p.font.size = Pt(11)
            p.font.color.rgb = rgb(SLATE)
            p.space_before = Pt(4)

    _apply_font(prs)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _apply_font(prs) -> None:
    def frames(shape):
        if shape.has_text_frame:
            yield shape.text_frame
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame

    for slide in prs.slides:
        for shape in slide.shapes:
            for tf in frames(shape):
                for p in tf.paragraphs:
                    p.font.name = FONT
                    for run in p.runs:
                        run.font.name = FONT


def _summary_slide(prs, blank, c: BriefContent, header, rect, text, bullets, Inches, Pt, rgb) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    slide = prs.slides.add_slide(blank)
    header(slide, "Impact brief", c.title)
    tf = text(slide, Inches(0.6), Inches(1.6), Inches(7.6), Inches(0.4), "Impact statement", 13, INDIGO, True)
    p = tf.add_paragraph()
    p.text = c.statement or c.summary or "Not generated yet."
    p.font.size = Pt(16)
    p.font.color.rgb = rgb(INK)
    p.space_before = Pt(6)
    if c.draft_note:
        box = rect(slide, Inches(0.6), Inches(5.6), Inches(7.6), Inches(1.0), AMBER_BG, MSO_SHAPE.ROUNDED_RECTANGLE)
        bt = box.text_frame
        bt.word_wrap = True
        bt.paragraphs[0].text = f"{c.draft_note} ({c.unconfirmed_count} unconfirmed)"
        bt.paragraphs[0].font.size = Pt(11)
        bt.paragraphs[0].font.color.rgb = rgb(AMBER)

    rect(slide, Inches(8.6), Inches(1.6), Inches(4.1), Inches(5.0), SURFACE, MSO_SHAPE.ROUNDED_RECTANGLE)
    y = 1.8
    for label, value in c.facts:
        text(slide, Inches(8.85), Inches(y), Inches(3.7), Inches(0.3), label.upper(), 9, MUTED, True)
        text(slide, Inches(8.85), Inches(y + 0.25), Inches(3.7), Inches(0.5), value, 14, INK, True)
        y += 0.95


def _effects_slide(prs, blank, c: BriefContent, header, text, bullets, Inches, Pt, rgb) -> None:
    slide = prs.slides.add_slide(blank)
    header(slide, "Expected effects", "Where the impact lands")
    rows = c.effects[:8]
    if not rows:
        text(slide, Inches(0.6), Inches(1.8), Inches(12), Inches(0.6), "No effects estimated yet.", 16, SLATE)
        return
    top = Inches(1.6)
    table = slide.shapes.add_table(len(rows) + 1, 6, Inches(0.6), top, Inches(12.1), Inches(0.5 * (len(rows) + 1))).table
    widths = [3.0, 1.6, 2.9, 1.5, 1.6, 1.5]
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    heads = ["Dimension", "Direction", "Range", "Confidence", "Status", "Sources"]
    for i, h in enumerate(heads):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(INK)
        cell.text = h
        para = cell.text_frame.paragraphs[0]
        para.font.size = Pt(12)
        para.font.bold = True
        para.font.color.rgb = rgb(WHITE)
    for r, e in enumerate(rows, start=1):
        values = [e.dimension, e.direction, e.range_text, e.confidence, "Confirmed" if e.confirmed else "Unconfirmed", e.citations or "—"]
        for i, v in enumerate(values):
            cell = table.cell(r, i)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(WHITE if r % 2 else SURFACE)
            cell.text = v
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(12)
            para.font.color.rgb = rgb(GREEN if (i == 4 and e.confirmed) else AMBER if i == 4 else INK)
            para.font.bold = i in (0, 4)
    if c.exploratory_note:
        text(slide, Inches(0.6), Inches(6.4), Inches(12), Inches(0.5), c.exploratory_note, 11, SLATE)
    if c.kpis:
        y = 1.6 + 0.5 * (len(rows) + 1) + 0.3
        if y < 5.8:
            tf = text(slide, Inches(0.6), Inches(y), Inches(12), Inches(0.4), "How we will measure it", 13, INDIGO, True)
            bullets(tf, [f"{n}: {d}" for n, d in c.kpis[:3]], 12)
