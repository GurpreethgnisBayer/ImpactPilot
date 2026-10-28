"""Quick brief -> compendium tests with mocked search and LLM."""

from __future__ import annotations

import json

import pytest
from fastapi.testclient import TestClient

from impactpilot.domain.models import InitiativeRecord, ValueSource
from impactpilot.domain.ranges import unconfirmed_llm_values
from impactpilot.main import app
from impactpilot.persistence.db import init_db
from impactpilot.services import compendium as comp
from impactpilot.services.llm_provider import OpenAICompatibleProvider
from impactpilot.services.openalex import _abstract_from_inverted_index
from impactpilot.services.sources import registry

from conftest import stub_sources

PUBMED = [
    {
        "pmid": "111",
        "title": "Automated HTS triage reduces cycle time",
        "abstract": "Automated triage reduced screening cycle time by 30% across 12 campaigns.",
        "journal": "J Screen",
        "year": "2023",
        "url": "https://pubmed.ncbi.nlm.nih.gov/111/",
    }
]
OPENALEX = [
    {
        "openalex_id": "W222",
        "title": "Machine learning for hit triage: an industry report",
        "abstract": "Teams reported 20-40% fewer manual review hours.",
        "year": 2024,
        "venue": "ChemRxiv",
        "type": "preprint",
        "url": "https://doi.org/10.1/x",
    }
]


def _fake_pubmed(**_kw):
    return PUBMED


def _fake_openalex(_q, **_kw):
    return OPENALEX


# Most cases only exercise PubMed and OpenAlex.
_JOURNALS_ONLY = dict(sources=["pubmed", "openalex"])


def _fake_llm(prompt: str) -> str:
    if "pubmed_query" in prompt:
        return '```json\n{"pubmed_query": "hit triage AND automation", "keyword_query": "ML hit triage"}\n```'
    return json.dumps(
        {
            "executive_summary": "Automated triage shortens screening cycles.",
            "value_chain_phase": "lead_generation_optimization",
            "process_step": "HTS hit triage",
            "mechanism": "ML ranks hits so chemists review fewer compounds.",
            "key_findings": [
                {"finding": "30% cycle-time reduction", "citations": ["PMID:111"]},
                {"finding": "Invented claim", "citations": ["PMID:999"]},
            ],
            "dimensions": [
                {
                    "dimension_id": "cycle_time",
                    "direction": "down",
                    "low": 40, "base": 20, "high": 30,
                    "unit": "% reduction",
                    "confidence": "medium",
                    "rationale": "Reported in automated triage studies.",
                    "citations": ["PMID:111", "OA:W222", "PMID:999"],
                    "quotes": ["reduced screening cycle time by 30%"],
                },
                {"dimension_id": "not_a_dimension", "low": 1, "base": 2, "high": 3},
            ],
            "suggested_maturity": "emerging_local",
            "maturity_rationale": "Published effects exist.",
            "suggested_kpis": [{"name": "Triage cycle time", "definition": "days", "cadence": "monthly"}],
            "risks_and_caveats": ["Model drift"],
            "open_questions": ["Current baseline?"],
        }
    )


def test_parse_llm_json_handles_fences_and_prose():
    assert comp.parse_llm_json('Sure!\n```json\n{"a": 1}\n```') == {"a": 1}
    assert comp.parse_llm_json('Here: {"a": 2} done') == {"a": 2}
    with pytest.raises(ValueError):
        comp.parse_llm_json("no json")


def test_openalex_abstract_reconstruction():
    assert _abstract_from_inverted_index({"world": [1], "hello": [0]}) == "hello world"


def test_openai_compatible_url_with_v1_suffix(monkeypatch):
    seen = {}

    class R:
        def raise_for_status(self):
            pass

        def json(self):
            return {"choices": [{"message": {"content": "ok"}}]}

    def fake_post(url, **_kw):
        seen["url"] = url
        return R()

    monkeypatch.setattr("impactpilot.services.llm_provider.requests.post", fake_post)
    OpenAICompatibleProvider("https://x.openai.azure.com/openai/v1", "k", "m").generate("hi")
    assert seen["url"] == "https://x.openai.azure.com/openai/v1/chat/completions"


def test_generate_compendium_with_llm_validates_citations():
    c = comp.generate_compendium(
        title="ML hit triage",
        description="Use ML to rank HTS hits for chemist review",
        generate_fn=_fake_llm,
        pubmed_fn=_fake_pubmed,
        openalex_fn=_fake_openalex,
        **_JOURNALS_ONLY,
    )
    assert c["llm_used"] is True
    assert c["queries"]["pubmed"] == "hit triage AND automation"
    assert {s["source_id"] for s in c["sources"]} == {"PMID:111", "OA:W222"}
    assert c["key_findings"][1]["citations"] == []
    assert len(c["dimensions"]) == 1
    d = c["dimensions"][0]
    assert (d["low"], d["base"], d["high"]) == (20, 30, 40)
    assert d["citations"] == ["PMID:111", "OA:W222"]
    assert c["value_chain_phase"] == "lead_generation_optimization"


def test_dimension_is_grounded_checks_low_and_high_against_source_numbers():
    sources_by_id = {
        "PMID:1": {"abstract": "Cycle time dropped from baseline by 20-40% across sites."},
        "PMID:2": {"abstract": "No numeric results were reported in this review."},
    }
    grounded = {"low": 20, "high": 40, "citations": ["PMID:1"]}
    assert comp._dimension_is_grounded(grounded, sources_by_id) is True

    invented = {"low": 5, "high": 95, "citations": ["PMID:1"]}
    assert comp._dimension_is_grounded(invented, sources_by_id) is False

    uncited = {"low": 20, "high": 40, "citations": []}
    assert comp._dimension_is_grounded(uncited, sources_by_id) is False

    no_numbers = {"low": 20, "high": 40, "citations": ["PMID:2"]}
    assert comp._dimension_is_grounded(no_numbers, sources_by_id) is False


def test_dimension_is_grounded_handles_fraction_vs_percent():
    sources_by_id = {"PMID:1": {"abstract": "Technician review time is reduced by nearly 50%."}}
    # Dimension expressed as a 0-1 fraction rather than a percent.
    dim = {"low": 0.5, "high": 0.75, "citations": ["PMID:1"]}
    assert comp._dimension_is_grounded(dim, sources_by_id) is True


def test_llm_synthesis_downgrades_confidence_for_ungrounded_high_confidence_claim():
    def _llm_invented_range(prompt: str) -> str:
        if "pubmed_query" in prompt:
            return '```json\n{"pubmed_query": "hit triage", "keyword_query": "hit triage"}\n```'
        return json.dumps(
            {
                "executive_summary": "summary",
                "value_chain_phase": "lead_generation_optimization",
                "process_step": "HTS hit triage",
                "mechanism": "mechanism",
                "key_findings": [],
                "dimensions": [
                    {
                        "dimension_id": "cycle_time",
                        "direction": "down",
                        # Neither edge appears anywhere near the source's reported "30%".
                        "low": 60, "base": 75, "high": 90,
                        "unit": "% reduction",
                        "confidence": "high",
                        "rationale": "Strongly supported.",
                        "citations": ["PMID:111"],
                        "quotes": ["reduced screening cycle time by 30%"],
                    }
                ],
                "suggested_maturity": "mature",
                "maturity_rationale": "",
                "suggested_kpis": [],
                "risks_and_caveats": [],
                "open_questions": [],
            }
        )

    c = comp.generate_compendium(
        title="ML hit triage",
        description="Use ML to rank HTS hits for chemist review",
        generate_fn=_llm_invented_range,
        pubmed_fn=_fake_pubmed,
        openalex_fn=lambda *_a, **_k: [],
        **_JOURNALS_ONLY,
    )
    d = c["dimensions"][0]
    assert d["grounded"] is False
    # The model claimed "high" confidence but nothing anchors the range to the source, so it's
    # capped at "medium" rather than feeding an unverified "high" into the impact score.
    assert d["confidence"] == "medium"


def test_generate_compendium_without_llm_lists_literature():
    c = comp.generate_compendium(
        title="ML hit triage",
        description="Use ML to rank HTS hits for chemist review",
        pubmed_fn=_fake_pubmed,
        openalex_fn=_fake_openalex,
        **_JOURNALS_ONLY,
    )
    assert c["llm_used"] is False
    assert len(c["sources"]) == 2
    assert c["dimensions"] == []
    assert any(n["raw"].startswith("30") for n in c["numeric_snippets"])


def test_off_topic_papers_are_dropped():
    hits = [
        {"pmid": "1", "title": "AI histopathology of preclinical toxicology slides", "abstract": "toxicology histopathology"},
        {"pmid": "2", "title": "Xpert assay for tuberculosis", "abstract": "automated clinical trial of rifampicin"},
    ]
    c = comp.generate_compendium(
        title="AI histopathology review for preclinical toxicology",
        description="Deep learning pre-screens tissue slides in preclinical safety studies",
        pubmed_fn=lambda **_k: hits,
        openalex_fn=lambda *_a, **_k: [],
        **_JOURNALS_ONLY,
    )
    assert [s["source_id"] for s in c["sources"]] == ["PMID:1"]


def test_search_failure_is_reported_not_raised():
    def boom(**_kw):
        raise RuntimeError("down")

    c = comp.generate_compendium(
        title="x idea",
        description="something long enough",
        pubmed_fn=boom,
        openalex_fn=_fake_openalex,
        **_JOURNALS_ONLY,
    )
    assert any("PubMed search failed" in w for w in c["warnings"])
    assert len(c["sources"]) == 1


def test_extra_sources_dedupe_and_keep_their_ids():
    europepmc = [
        {
            "source_id": "PMID:111",
            "database": "europepmc",
            "pmid": "111",
            "title": "Another writeup of HTS triage",
            "abstract": "Same PubMed record, should be dropped.",
        },
        {
            "source_id": "PMC:PMC999",
            "database": "europepmc",
            "pmid": "999",
            "title": "NHS guideline on HTS triage",
            "abstract": "A public report on triage.",
        },
    ]
    trial = {
        "source_id": "NCT:NCT1",
        "database": "clinicaltrials",
        "title": "HTS triage operational trial",
        "abstract": "Enrollment and cycle time.",
        "type": "trial",
    }
    label = {
        "source_id": "FDA:abc",
        "database": "openfda",
        "title": "HTS triage drug label",
        "abstract": "Indications for triage.",
        "type": "label",
    }
    scholar = [
        {
            "source_id": "S2:dup",
            "database": "semanticscholar",
            "doi": "10.1/x",
            "title": "DOI twin of the HTS triage report",
            "abstract": "triage",
        },
        {
            "source_id": "S2:keep",
            "database": "semanticscholar",
            "doi": "10.2/keep",
            "title": "Industry report on HTS triage",
            "abstract": "triage methods",
        },
    ]
    openalex = [{**OPENALEX[0], "doi": "https://doi.org/10.1/x"}]
    c = comp.generate_compendium(
        title="ML hit triage",
        description="Use ML to rank HTS hits for chemist review",
        pubmed_fn=_fake_pubmed,
        openalex_fn=lambda *_a, **_k: openalex,
        europepmc_fn=lambda *_a, **_k: europepmc,
        clinicaltrials_fn=lambda *_a, **_k: [trial],
        openfda_fn=lambda *_a, **_k: [label],
        semanticscholar_fn=lambda *_a, **_k: scholar,
    )
    by_id = {s["source_id"]: s for s in c["sources"]}
    assert "PMID:111" in by_id and by_id["PMID:111"]["database"] == "pubmed"
    assert by_id["PMC:PMC999"]["database"] == "europepmc"
    assert "S2:dup" not in by_id
    assert by_id["NCT:NCT1"]["database"] == "clinicaltrials"
    assert by_id["FDA:abc"]["database"] == "openfda"
    assert by_id["S2:keep"]["database"] == "semanticscholar"
    assert by_id["OA:W222"]["database"] == "openalex"
    assert set(c["queries"]) >= {"pubmed", "keyword", "europepmc", "clinicaltrials", "openfda", "semanticscholar"}


def test_failed_source_does_not_block_the_others():
    def boom(*_a, **_k):
        raise RuntimeError("rate limited")

    c = comp.generate_compendium(
        title="ML hit triage",
        description="Use ML to rank HTS hits for chemist review",
        pubmed_fn=_fake_pubmed,
        openalex_fn=_fake_openalex,
        include_europepmc=False,
        include_clinicaltrials=False,
        include_semanticscholar=False,
        openfda_fn=boom,
    )
    assert {s["source_id"] for s in c["sources"]} == {"PMID:111", "OA:W222"}
    assert any("openFDA drug labels search failed" in w and "rate limited" in w for w in c["warnings"])


def test_group_quota_keeps_trials_and_regulatory_when_literature_is_heavy():
    papers = [
        {"pmid": str(i), "title": f"HTS hit triage machine learning study {i}", "abstract": "hit triage HTS ranking"}
        for i in range(1, 60)
    ]
    trial = {"source_id": "NCT:NCT1", "title": "HTS triage operations", "abstract": "triage workflow", "type": "trial"}
    label = {"source_id": "FDA:x", "title": "Triage software label", "abstract": "HTS triage", "type": "label"}
    c = comp.generate_compendium(
        title="ML hit triage",
        description="Use ML to rank HTS hits for chemist review",
        sources=["pubmed", "clinicaltrials", "openfda"],
        overrides={
            "pubmed": lambda _q, _n: papers,
            "clinicaltrials": lambda _q, _n: [trial],
            "openfda": lambda _q, _n: [label],
        },
    )
    ids = [s["source_id"] for s in c["sources"]]
    assert len(ids) == comp.MAX_MERGED_SOURCES
    assert "NCT:NCT1" in ids and "FDA:x" in ids


def test_source_needing_a_key_is_skipped_with_a_warning(monkeypatch):
    import dataclasses

    spec = dataclasses.replace(registry.SOURCES["openfda"], env_keys=("FAKE_SOURCE_KEY",))
    monkeypatch.setitem(registry.SOURCES, "openfda", spec)
    monkeypatch.delenv("FAKE_SOURCE_KEY", raising=False)
    assert "openfda" not in registry.default_keys()
    c = comp.generate_compendium(
        title="ML hit triage",
        description="Use ML to rank HTS hits for chemist review",
        sources=["pubmed", "openfda"],
        overrides={"pubmed": lambda _q, _n: PUBMED},
    )
    assert any("set FAKE_SOURCE_KEY" in w for w in c["warnings"])
    monkeypatch.setenv("FAKE_SOURCE_KEY", "x")
    assert "openfda" in registry.default_keys()


def test_slow_source_times_out_without_blocking(monkeypatch):
    import time

    monkeypatch.setattr(comp, "SOURCE_TIMEOUT_S", 0.2)
    c = comp.generate_compendium(
        title="ML hit triage",
        description="Use ML to rank HTS hits for chemist review",
        sources=["pubmed", "openfda"],
        overrides={"pubmed": lambda _q, _n: PUBMED, "openfda": lambda _q, _n: time.sleep(1) or []},
    )
    assert [s["source_id"] for s in c["sources"]] == ["PMID:111"]
    assert any("timed out" in w for w in c["warnings"])


def test_unknown_legacy_argument_is_rejected():
    with pytest.raises(TypeError, match="include_nope"):
        comp.generate_compendium(title="x idea", description="long enough text", include_nope=True)


def test_source_clients_parse_payloads(monkeypatch):
    class Response:
        status_code = 200

        def __init__(self, payload):
            self._payload = payload

        def raise_for_status(self):
            return None

        def json(self):
            return self._payload

    def route(url, **_kwargs):
        if "europepmc" in url:
            return Response(
                {
                    "resultList": {
                        "result": [
                            {
                                "pmid": "5",
                                "pmcid": "PMC5",
                                "doi": "10.1/e",
                                "title": "A guideline",
                                "abstractText": "Public guidance.",
                                "pubYear": "2021",
                                "journalTitle": "NICE",
                            }
                        ]
                    }
                }
            )
        if "clinicaltrials" in url:
            return Response(
                {
                    "studies": [
                        {
                            "hasResults": True,
                            "protocolSection": {
                                "identificationModule": {"nctId": "NCT9", "officialTitle": "A trial"},
                                "descriptionModule": {"briefSummary": "Patients enrolled."},
                                "statusModule": {"startDateStruct": {"date": "2022-01"}},
                                "sponsorCollaboratorsModule": {"leadSponsor": {"name": "Sponsor"}},
                            },
                        }
                    ]
                }
            )
        if "fda.gov" in url:
            return Response(
                {
                    "results": [
                        {
                            "set_id": "abc",
                            "effective_time": "20200301",
                            "openfda": {"brand_name": ["DrugX"], "manufacturer_name": ["Maker"]},
                            "indications_and_usage": ["Treats the condition."],
                        }
                    ]
                }
            )
        return Response(
            {
                "data": [
                    {
                        "paperId": "p1",
                        "title": "A report",
                        "abstract": "Findings.",
                        "year": 2024,
                        "venue": "Working paper",
                        "externalIds": {"DOI": "10.2/s", "PubMed": "8"},
                        "url": "https://www.semanticscholar.org/paper/p1",
                    }
                ]
            }
        )

    for module in (
        "impactpilot.services.europepmc",
        "impactpilot.services.clinicaltrials",
        "impactpilot.services.openfda",
        "impactpilot.services.semantic_scholar",
    ):
        monkeypatch.setattr(f"{module}.requests.get", route)

    from impactpilot.services.clinicaltrials import search_clinicaltrials
    from impactpilot.services.europepmc import search_europepmc
    from impactpilot.services.openfda import search_openfda
    from impactpilot.services.semantic_scholar import search_semanticscholar

    epmc = search_europepmc("triage")
    assert epmc[0]["source_id"] == "PMC:PMC5" and epmc[0]["database"] == "europepmc" and epmc[0]["pmid"] == "5"
    trial = search_clinicaltrials("triage")
    assert trial[0]["source_id"] == "NCT:NCT9" and trial[0]["type"] == "trial"
    assert "Results have been posted" in trial[0]["abstract"]
    label = search_openfda("triage")
    assert label[0]["source_id"] == "FDA:abc" and label[0]["type"] == "label" and label[0]["year"] == "2020"
    paper = search_semanticscholar("triage")
    assert paper[0]["source_id"] == "S2:p1" and paper[0]["pmid"] == "8" and paper[0]["doi"] == "10.2/s"

    def limited(url, **_kwargs):
        class Denied:
            status_code = 429

            def raise_for_status(self):
                raise AssertionError("429 should be reported before raise_for_status")

            def json(self):
                return {}

        return Denied()

    monkeypatch.setattr("impactpilot.services.semantic_scholar.requests.get", limited)
    with pytest.raises(RuntimeError, match="429"):
        search_semanticscholar("triage")


def test_clinicaltrials_falls_back_to_strongest_terms(monkeypatch):
    seen: list[dict] = []

    class Response:
        status_code = 200

        def __init__(self, studies):
            self._studies = studies

        def raise_for_status(self):
            return None

        def json(self):
            return {"studies": self._studies}

    study = {
        "protocolSection": {
            "identificationModule": {"nctId": "NCT7", "briefTitle": "Home nursing visits"},
            "descriptionModule": {"briefSummary": "Remote visits."},
        }
    }

    def fake_get(_url, params, **_kw):
        seen.append(params)
        return Response([study] if "query.cond" in params else [])

    monkeypatch.setattr("impactpilot.services.clinicaltrials.requests.get", fake_get)
    from impactpilot.services.clinicaltrials import search_clinicaltrials

    hits = search_clinicaltrials("decentralized clinical trial home nursing")
    assert [h["source_id"] for h in hits] == ["NCT:NCT7"]
    assert [next(k for k in p if k.startswith("query.")) for p in seen] == ["query.term", "query.intr", "query.cond"]
    assert seen[1]["query.intr"] == "decentralized nursing"


def _fake_curation_llm(prompt: str) -> str:
    if "pubmed_query" in prompt:
        return '{"pubmed_query": "hit triage AND automation", "keyword_query": "ML hit triage"}'
    if "recommended_ids" in prompt:
        return json.dumps(
            {
                "interpretation": "This reads as automating HTS hit triage with ML to cut chemist review time.",
                "recommended_ids": ["PMID:111"],
                "notes": {"PMID:111": "Directly about automated triage cycle time.", "OA:W222": "Different domain."},
            }
        )
    return json.dumps({"executive_summary": "x", "dimensions": []})


def test_search_literature_flags_recommended_sources():
    result = comp.search_literature(
        title="ML hit triage",
        description="Use ML to rank HTS hits for chemist review",
        generate_fn=_fake_curation_llm,
        pubmed_fn=_fake_pubmed,
        openalex_fn=_fake_openalex,
        **_JOURNALS_ONLY,
    )
    assert result["llm_used"] is True
    assert "automating HTS hit triage" in result["interpretation"]
    by_id = {s["source_id"]: s for s in result["sources"]}
    assert by_id["PMID:111"]["recommended"] is True
    assert by_id["OA:W222"]["recommended"] is False
    assert by_id["OA:W222"]["note"] == "Different domain."
    assert result["queries"]["pubmed"] == "hit triage AND automation"


def test_search_literature_without_llm_recommends_all_filtered():
    result = comp.search_literature(
        title="ML hit triage",
        description="Use ML to rank HTS hits for chemist review",
        pubmed_fn=_fake_pubmed,
        openalex_fn=_fake_openalex,
        **_JOURNALS_ONLY,
    )
    assert result["llm_used"] is False
    assert result["interpretation"] == ""
    assert all(s["recommended"] for s in result["sources"])


def test_synthesize_orders_sources_by_user_relevance():
    seen: dict[str, str] = {}

    def _capture(prompt: str) -> str:
        seen["prompt"] = prompt
        return _fake_llm(prompt)

    c = comp.synthesize_from_sources(
        title="ML hit triage",
        description="Use ML to rank HTS hits for chemist review",
        sources=[
            {
                "source_id": "OA:W222",
                "database": "openalex",
                "title": OPENALEX[0]["title"],
                "abstract": OPENALEX[0]["abstract"],
                "relevance": 1,
            },
            {
                "source_id": "PMID:111",
                "database": "pubmed",
                "title": PUBMED[0]["title"],
                "abstract": PUBMED[0]["abstract"],
                "relevance": "high",
            },
            {
                "source_id": "PMID:999",
                "database": "pubmed",
                "title": "Off topic paper",
                "abstract": "Should be dropped.",
                "relevance": 0,
            },
        ],
        generate_fn=_capture,
    )
    assert [s["source_id"] for s in c["sources"]] == ["PMID:111", "OA:W222"]
    assert [s["relevance"] for s in c["sources"]] == ["high", "low"]
    block = seen["prompt"].split("Sources (cite ONLY these ids):", 1)[1].split("Numbers found", 1)[0]
    assert block.index("User relevance: high") < block.index("User relevance: low")
    assert "PMID:999" not in block
    assert "must not override a higher-ranked source" in seen["prompt"]


def test_synthesize_rejects_when_every_source_is_off():
    with pytest.raises(ValueError, match="Rank at least one"):
        comp.synthesize_from_sources(
            title="ML hit triage",
            description="Use ML to rank HTS hits for chemist review",
            sources=[{"source_id": "PMID:111", "title": "x", "abstract": "y", "relevance": "off"}],
            generate_fn=_fake_llm,
        )


def test_synthesize_from_sources_uses_given_sources_only():
    chosen = [
        {"source_id": "PMID:111", "database": "pubmed", "title": PUBMED[0]["title"], "abstract": PUBMED[0]["abstract"], "year": "2023"},
    ]
    c = comp.synthesize_from_sources(
        title="ML hit triage",
        description="Use ML to rank HTS hits for chemist review",
        sources=chosen,
        generate_fn=_fake_llm,
    )
    assert c["llm_used"] is True
    assert [s["source_id"] for s in c["sources"]] == ["PMID:111"]
    # The fake LLM cites OA:W222 and a made-up PMID too, but only PMID:111 was in scope.
    assert c["dimensions"][0]["citations"] == ["PMID:111"]


def test_apply_compendium_marks_values_unconfirmed():
    c = comp.generate_compendium(
        title="ML hit triage",
        description="Use ML to rank HTS hits",
        generate_fn=_fake_llm,
        pubmed_fn=_fake_pubmed,
        openalex_fn=_fake_openalex,
        **_JOURNALS_ONLY,
    )
    rec = comp.apply_compendium(InitiativeRecord(), c, title="ML hit triage", description="d")
    assert rec.scope.phase_id == "lead_generation_optimization"
    assert rec.maturity.tier.value == "emerging_local"
    est = rec.estimates[0].estimate
    assert est.source == ValueSource.llm_suggested and est.confirmed is False
    assert len(est.evidence_ids) == 2
    assert unconfirmed_llm_values(rec)
    assert rec.measurement_plan[0].name == "Triage cycle time"


@pytest.fixture()
def client(tmp_path, monkeypatch):
    db_url = f"sqlite:///{tmp_path / 'test.db'}"
    monkeypatch.setenv("IMPACTPILOT_DB_URL", db_url)
    for k in ("OPENAI_API_KEY", "AZURE_OPENAI_API_KEY"):
        monkeypatch.delenv(k, raising=False)
    stub_sources(monkeypatch, pubmed=lambda _q, _n: PUBMED, openalex=lambda _q, _n: OPENALEX)
    monkeypatch.setattr(comp, "build_generate_fn", lambda _s, timeout=180: _fake_llm)
    import impactpilot.persistence.db as dbmod

    dbmod._ENGINE = None
    dbmod._SessionLocal = None
    init_db(db_url)
    with TestClient(app) as c:
        yield c


def test_compendium_api_flow(client):
    llm = {"provider": "openai_compatible", "openai_base_url": "https://x.openai.azure.com/openai/v1",
           "openai_api_key": "k", "openai_model": "m"}
    r = client.post(
        "/api/compendium",
        json={"title": "ML hit triage", "description": "Use ML to rank HTS hits for review", "llm_settings": llm},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    iid = body["initiative_id"]
    assert body["compendium"]["llm_used"] is True

    md = client.get(f"/api/initiatives/{iid}/compendium.md")
    assert md.status_code == 200
    assert "Impact Compendium" in md.text and "[1]" in md.text

    draft = client.get(f"/api/initiatives/{iid}/export.md")
    assert draft.status_code == 200
    assert "DRAFT" in draft.text and "**unconfirmed**" in draft.text
    assert "-draft.md" in draft.headers["content-disposition"]
    assert "Key findings" in draft.text and "References" in draft.text

    status = client.get(f"/api/initiatives/{iid}/export-status").json()
    assert status["status"] == "draft" and status["unconfirmed_count"] > 0

    pdf = client.get(f"/api/initiatives/{iid}/export.pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF") and len(pdf.content) > 2000
    pptx = client.get(f"/api/initiatives/{iid}/export.pptx")
    assert pptx.status_code == 200 and pptx.content[:2] == b"PK"
    import io

    from pptx import Presentation

    deck = Presentation(io.BytesIO(pptx.content))
    texts = " ".join(sh.text_frame.text for s in deck.slides for sh in s.shapes if sh.has_text_frame)
    assert len(deck.slides) >= 3 and "DRAFT" in texts and "References" in texts

    r = client.post(f"/api/initiatives/{iid}/estimates/cycle_time/confirm", json={"base": 25})
    assert r.status_code == 200
    est = r.json()["estimates"][0]["estimate"]
    assert est["confirmed"] is True and est["source"] == "user" and est["base"] == 25

    assert client.post(f"/api/initiatives/{iid}/estimates/confirm-all").status_code == 200
    final = client.get(f"/api/initiatives/{iid}/export.md")
    assert final.status_code == 200 and "DRAFT" not in final.text
    assert client.get(f"/api/initiatives/{iid}/export-status").json()["status"] == "final"


def test_guided_flow_search_then_synthesize_endpoints(client):
    llm = {"provider": "openai_compatible", "openai_base_url": "https://x.openai.azure.com/openai/v1",
           "openai_api_key": "k", "openai_model": "m"}
    r = client.post(
        "/api/compendium/search",
        json={"title": "ML hit triage", "description": "Use ML to rank HTS hits for chemist review", "llm_settings": llm},
    )
    assert r.status_code == 200, r.text
    search = r.json()
    assert len(search["sources"]) == 2
    kept = [s for s in search["sources"] if s["recommended"]]
    assert kept

    r = client.post(
        "/api/compendium/synthesize",
        json={
            "title": "ML hit triage",
            "description": "Use ML to rank HTS hits for chemist review",
            "sources": kept,
            "llm_settings": llm,
        },
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["compendium"]["llm_used"] is True
    assert {s["source_id"] for s in body["compendium"]["sources"]} == {s["source_id"] for s in kept}

    r = client.post("/api/compendium/synthesize", json={"title": "x", "description": "y" * 10, "sources": []})
    assert r.status_code == 422


def test_effective_settings_ignores_incomplete_override(monkeypatch):
    monkeypatch.setenv("IMPACTPILOT_LLM_PROVIDER", "openai_compatible")
    monkeypatch.setenv("OPENAI_BASE_URL", "https://example.openai.azure.com/openai/v1")
    monkeypatch.setenv("OPENAI_API_KEY", "server-key")
    monkeypatch.setenv("OPENAI_MODEL", "gpt-4o")
    from impactpilot.api.compendium import _effective_settings

    # A stale override for a removed local provider must not break the platform default.
    resolved = _effective_settings({"provider": "ollama", "ollama_host": "http://localhost:11434"})
    assert resolved is not None
    assert resolved["provider"] == "openai_compatible"
    assert resolved["openai_api_key"] == "server-key"
    assert "ollama_host" not in resolved

    # A partial override of a supported provider also falls back to the platform default.
    resolved = _effective_settings({"provider": "azure_openai"})
    assert resolved is not None
    assert resolved["provider"] == "openai_compatible"


def test_resolve_llm_settings_keeps_server_key(monkeypatch):
    monkeypatch.setenv("IMPACTPILOT_LLM_PROVIDER", "openai_compatible")
    monkeypatch.setenv("OPENAI_BASE_URL", "https://example.openai.azure.com/openai/v1")
    monkeypatch.setenv("OPENAI_API_KEY", "server-key")
    monkeypatch.setenv("OPENAI_MODEL", "gpt-4o")
    from impactpilot.api.llm import resolve_llm_settings

    resolved = resolve_llm_settings({"openai_model": "o4-mini", "openai_api_key": "", "provider": "openai_compatible"})
    assert resolved is not None
    assert resolved["openai_api_key"] == "server-key"
    assert resolved["openai_model"] == "o4-mini"
    assert resolve_llm_settings(None)["openai_model"] == "gpt-4o"
    monkeypatch.delenv("OPENAI_API_KEY")
    assert resolve_llm_settings(None) is None


def test_compendium_api_without_llm(client):
    r = client.post("/api/compendium", json={"title": "ML hit triage", "description": "Use ML to rank HTS hits"})
    assert r.status_code == 200
    assert r.json()["compendium"]["llm_used"] is False


def test_sources_endpoint_lists_every_registered_source(client):
    body = client.get("/api/compendium/sources").json()
    keys = [s["key"] for s in body]
    assert keys == list(registry.SOURCES)
    assert {"key", "label", "group", "group_label", "available", "needs", "default_on"} <= set(body[0])


def test_search_api_accepts_sources_list_and_legacy_flags(client):
    r = client.post(
        "/api/compendium/search",
        json={"title": "ML hit triage", "description": "Use ML to rank HTS hits", "sources": ["openalex"]},
    )
    assert {s["source_id"] for s in r.json()["sources"]} == {"OA:W222"}
    legacy = client.post(
        "/api/compendium/search",
        json={"title": "ML hit triage", "description": "Use ML to rank HTS hits", "include_openalex": False},
    )
    assert {s["source_id"] for s in legacy.json()["sources"]} == {"PMID:111"}
